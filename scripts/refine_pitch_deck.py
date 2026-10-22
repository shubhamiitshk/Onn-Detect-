import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

sys.stdout.reconfigure(encoding='utf-8')

src_path = r'C:\Users\PLEXTEK\.gemini\antigravity\brain\ab162997-2a8d-4152-a5b1-906b2f5331ee\.user_uploaded\media_1792663862645_69ab35af.pptx'
dst_path_root = r'c:\Users\PLEXTEK\Desktop\Onn-Detect-\ONN_DETECT_PITCH_DECK_Z21_KGP_FORGE.pptx'
dst_path_docs = r'c:\Users\PLEXTEK\Desktop\Onn-Detect-\docs\ONN_DETECT_PITCH_DECK_Z21_KGP_FORGE.pptx'
dst_path_artifact = r'C:\Users\PLEXTEK\.gemini\antigravity\brain\ab162997-2a8d-4152-a5b1-906b2f5331ee\ONN_DETECT_PITCH_DECK_Z21_KGP_FORGE.pptx'

prs = Presentation(src_path)

# --- SLIDE 1 ---
# Shape 1 P2: Under 450 ms -> Under 180 ms
s1_sh1 = prs.slides[0].shapes[1]
if len(s1_sh1.text_frame.paragraphs) >= 3:
    p = s1_sh1.text_frame.paragraphs[2]
    if len(p.runs) > 0:
        p.runs[0].text = "On-device deepfake, voice-clone, and cyber-fraud defense in under 180 ms."

# Shape 2 P0: Metadata
s1_sh2 = prs.slides[0].shapes[2]
if len(s1_sh2.text_frame.paragraphs) >= 1:
    p = s1_sh2.text_frame.paragraphs[0]
    if len(p.runs) > 0:
        p.runs[0].text = "Domain: AI & Deep Tech  |  Stage: Pre-Seed (TRL 6-7 Validated Prototype)  |  IIT Kharagpur"

# --- SLIDE 3 ---
# Shape 3 P1: Sub-450 ms -> Sub-180 ms & C2PA + AudioSeal
s3_sh3 = prs.slides[2].shapes[3]
if len(s3_sh3.text_frame.paragraphs) >= 2:
    p1 = s3_sh3.text_frame.paragraphs[1]
    if len(p1.runs) >= 2:
        p1.runs[0].text = "Sub-180 ms Edge Cascade: "
        p1.runs[1].text = "Unifies INT8 SigLIP-B/16, 36-ch SRM/NPR CNN, LCNN-MFM speech anti-spoofing (<25 ms), and C2PA & AudioSeal verification."

# --- SLIDE 4 ---
# Shape 7 P0: AI Reel Auto-Skipper & Instant Guard
s4_sh7 = prs.slides[3].shapes[7]
if len(s4_sh7.text_frame.paragraphs) >= 4:
    p0 = s4_sh7.text_frame.paragraphs[0]
    if len(p0.runs) >= 1:
        p0.runs[0].text = "AI Reel Filter & Forward Guard"
    p2 = s4_sh7.text_frame.paragraphs[2]
    if len(p2.runs) >= 2:
        p2.runs[0].text = "Auto-Swipe & HUD: "
        p2.runs[1].text = "Swipes past AI Reels or displays a live corner authenticity badge."
    p3 = s4_sh7.text_frame.paragraphs[3]
    if len(p3.runs) >= 2:
        p3.runs[0].text = "WhatsApp/Share Guard: "
        p3.runs[1].text = "1-tap Share Sheet check on forwarded voice notes, viral videos, and QR scams."

# --- SLIDE 6 ---
# Shape 4 P3: Categorical C2PA -> Categorical Provenance (C2PA + AudioSeal)
s6_sh4 = prs.slides[5].shapes[4]
if len(s6_sh4.text_frame.paragraphs) >= 4:
    p3 = s6_sh4.text_frame.paragraphs[3]
    if len(p3.runs) >= 2:
        p3.runs[0].text = "Categorical Provenance: "
        p3.runs[1].text = "Treats cryptographic manifests (C2PA) and neural watermarks (AudioSeal) as positive proof."

# --- SLIDE 8 ---
# Table Row 4 Col 1 & Col 2: Latency & Row 5: StrongBox TEE
s8_table = prs.slides[7].shapes[2].table
# Row 4 Col 1: Yes (<25 ms Audio / <180 ms)
cell_r4c1 = s8_table.rows[4].cells[1]
if len(cell_r4c1.text_frame.paragraphs) > 0 and len(cell_r4c1.text_frame.paragraphs[0].runs) > 0:
    cell_r4c1.text_frame.paragraphs[0].runs[0].text = "Yes (<25 ms Audio / <180 ms)"
else:
    cell_r4c1.text = "Yes (<25 ms Audio / <180 ms)"

cell_r4c2 = s8_table.rows[4].cells[2]
if len(cell_r4c2.text_frame.paragraphs) > 0 and len(cell_r4c2.text_frame.paragraphs[0].runs) > 0:
    cell_r4c2.text_frame.paragraphs[0].runs[0].text = "Impossible (1-3s Network Lag)"
else:
    cell_r4c2.text = "Impossible (1-3s Network Lag)"

# Row 5 Col 0, 1, 2, 3: Update to include StrongBox TEE
cell_r5c0 = s8_table.rows[5].cells[0]
if len(cell_r5c0.text_frame.paragraphs) > 0 and len(cell_r5c0.text_frame.paragraphs[0].runs) > 0:
    cell_r5c0.text_frame.paragraphs[0].runs[0].text = "Readiness & Provenance"
cell_r5c1 = s8_table.rows[5].cells[1]
if len(cell_r5c1.text_frame.paragraphs) > 0 and len(cell_r5c1.text_frame.paragraphs[0].runs) > 0:
    cell_r5c1.text_frame.paragraphs[0].runs[0].text = "TRL 6-7 | StrongBox TEE | IP Pipeline"
cell_r5c2 = s8_table.rows[5].cells[2]
if len(cell_r5c2.text_frame.paragraphs) > 0 and len(cell_r5c2.text_frame.paragraphs[0].runs) > 0:
    cell_r5c2.text_frame.paragraphs[0].runs[0].text = "Cloud SaaS | Server TLS only"
cell_r5c3 = s8_table.rows[5].cells[3]
if len(cell_r5c3.text_frame.paragraphs) > 0 and len(cell_r5c3.text_frame.paragraphs[0].runs) > 0:
    cell_r5c3.text_frame.paragraphs[0].runs[0].text = "TRL 2-3 Lab Papers | None"

# --- SLIDE 11 ---
# Add contact info and repository URL to Slide 11
s11_sh1 = prs.slides[10].shapes[1]
tf = s11_sh1.text_frame
p_contact = tf.add_paragraph()
p_contact.text = "Founder: Shubham Kumar (IIT Kharagpur)  |  Repo: https://github.com/shubhamiitshk/Onn-Detect-"
p_contact.font.size = Pt(14)
p_contact.font.color.rgb = RGBColor(148, 163, 184)

# Save to all locations
prs.save(dst_path_root)
print(f"Saved: {dst_path_root}")
prs.save(dst_path_docs)
print(f"Saved: {dst_path_docs}")
prs.save(dst_path_artifact)
print(f"Saved: {dst_path_artifact}")

# Also overwrite user uploaded file so user downloading that specific file gets the refined version!
prs.save(src_path)
print(f"Updated user uploaded source: {src_path}")
